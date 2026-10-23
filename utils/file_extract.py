"""
Extracts raw text from an uploaded resume file: .pdf, .docx, or .txt.
Works directly on the in-memory upload stream (no disk write needed),
which also makes it safe for batch/concurrent uploads with duplicate
filenames.
"""

import io
import os


def extract_text_from_storage(file_storage) -> str:
    """file_storage: a werkzeug FileStorage object (request.files[...])."""
    filename = file_storage.filename
    ext = os.path.splitext(filename)[1].lower()
    raw_bytes = file_storage.read()

    if ext == ".txt":
        return raw_bytes.decode("utf-8", errors="ignore")

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(raw_bytes))
        text_parts = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(text_parts)

    if ext == ".docx":
        import docx
        doc = docx.Document(io.BytesIO(raw_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    raise ValueError(f"Unsupported file type: {ext}. Please upload .pdf, .docx, or .txt")


# Kept for backward compatibility (path-based extraction, e.g. for scripts/tests)
def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".txt":
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        reader = PdfReader(filepath)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == ".docx":
        import docx
        doc = docx.Document(filepath)
        return "\n".join(p.text for p in doc.paragraphs)

    raise ValueError(f"Unsupported file type: {ext}. Please upload .pdf, .docx, or .txt")
